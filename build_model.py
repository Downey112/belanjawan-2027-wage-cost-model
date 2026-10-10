import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]

FONT = "Arial"
BLUE = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10, color="000000")
GREEN = Font(name=FONT, size=10, color="008000")
BOLD = Font(name=FONT, size=10, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="595959")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
YELLOW = PatternFill("solid", fgColor="FFFF00")
FLAG_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
TOP = Border(top=Side(style="thin", color="000000"))
RM = '#,##0.00;(#,##0.00);"-"'
RM0 = '#,##0;(#,##0);"-"'
PCT = '0.00%;(0.00%);"-"'
USER = "Provided by user (Belanjawan 2027 brief)"

wb = Workbook()


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, bold=False, align=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, size=font.size, bold=bold or font.bold,
                  italic=font.italic, color=font.color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c


def header_row(ws, row, labels, start_col=1):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=label)
        c.font = HEAD
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def section(ws, row, label, last_col):
    for col in range(1, last_col + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    put(ws, f"A{row}", label, BOLD)


# --------------------------------------------------------------- Assumptions
wa = wb.active
wa.title = "Assumptions"
put(wa, "A1", "Assumptions - Belanjawan 2027 Wage Parameters", TITLE)
put(wa, "A2", "Blue = editable input. Every other sheet reads from here. "
              "Orange rows are figures worth verifying against the official schedule.", NOTE)
header_row(wa, 4, ["Parameter", "Value", "Unit", "Source / note"])

A = {}  # name -> absolute reference on Assumptions
row = 5


def a_section(label):
    global row
    section(wa, row, label, 4)
    row += 1


def a_input(name, label, value, unit, note, fmt=RM, flag=False):
    global row
    put(wa, f"A{row}", label)
    put(wa, f"B{row}", value, BLUE, fmt)
    put(wa, f"C{row}", unit)
    put(wa, f"D{row}", note, NOTE, align=Alignment(wrap_text=True, vertical="top"))
    if flag:
        for col in "ABCD":
            wa[f"{col}{row}"].fill = FLAG_FILL
    A[name] = f"Assumptions!$B${row}"
    row += 1


a_section("Wage benchmarks")
a_input("minwage", "National Minimum Wage", 2000, "RM / month",
        USER + ". Effective June 2027; MSMEs under RM50m turnover temporarily exempt.")
a_input("fresh", "Fresh Graduate & Semi-Skilled Floor", 2500, "RM / month", USER + ".")
a_input("glc", "GLC / GLIC Living Wage Benchmark", 3400, "RM / month", USER + ".")

a_section("Employer on-costs")
a_input("er_epf_lo", "EPF employer rate (wage at or below threshold)", 0.13, "% of wage", USER + ".", PCT)
a_input("er_epf_hi", "EPF employer rate (wage above threshold)", 0.12, "% of wage",
        "Model assumption, not in the brief: current KWSP Third Schedule rate for wages above RM5,000.",
        PCT, flag=True)
a_input("epf_thr", "EPF rate threshold", 5000, "RM / month", USER + ".")
a_input("er_socso", "SOCSO employer rate (Employment Injury & Invalidity)", 0.0175, "% of wage",
        USER + " (stated as approx. 1.75%).", PCT)
a_input("er_socso_cap", "SOCSO employer cap", 103.65, "RM / month",
        USER + ". Check: 1.75% x RM6,000 = RM105.00 and the current PERKESO table shows RM104.15 "
               "at the RM6,000 ceiling, so RM103.65 looks slightly low.", flag=True)
a_input("er_eis", "EIS employer rate", 0.002, "% of wage", USER + ".", PCT)
a_input("er_eis_cap", "EIS employer cap", 11.90, "RM / month", USER + ".")
a_input("ceiling", "SOCSO / EIS insured wage ceiling", 6000, "RM / month", USER + ".")
a_input("hrd", "HRD Corp levy rate", 0.01, "% of basic + fixed allowances", USER + ".", PCT)

a_section("Employee deductions")
a_input("ee_epf", "EPF employee rate", 0.11, "% of wage", USER + ".", PCT)
a_input("ee_socso", "SOCSO employee rate", 0.005, "% of wage", USER + " (stated as approx. 0.50%).", PCT)
a_input("ee_socso_cap", "SOCSO employee cap", 29.65, "RM / month",
        USER + ". Check: 0.5% x RM6,000 = RM30.00 and the current PERKESO table shows RM29.75 "
               "at the ceiling.", flag=True)
a_input("ee_eis", "EIS employee rate", 0.002, "% of wage", USER + ".", PCT)
a_input("ee_eis_cap", "EIS employee cap", 11.90, "RM / month", USER + ".")

a_section("Income tax (PCB / MTD)")
a_input("relief_ind", "Individual relief", 12000, "RM / year",
        USER + ". Current statutory individual relief is RM9,000, so RM12,000 is a Belanjawan 2027 "
               "figure to confirm against the Finance Bill.", RM0, flag=True)
a_input("relief_epf", "EPF relief cap", 4000, "RM / year",
        "Model assumption, not in the brief: current LHDN cap on EPF contributions relief.", RM0, flag=True)
a_input("relief_other", "Other reliefs (per employee)", 0, "RM / year",
        "Model assumption: zero. Add lifestyle, SOCSO, medical etc. here if needed.", RM0)
a_input("rebate", "Individual tax rebate", 400, "RM / year", USER + ".", RM0)
a_input("rebate_lim", "Rebate chargeable-income limit", 35000, "RM / year",
        "Model assumption, not in the brief: current LHDN limit for the RM400 rebate.", RM0, flag=True)

row += 1
BR_HEAD = row
put(wa, f"A{row}", "Resident individual tax brackets", BOLD)
put(wa, f"D{row}", "Model assumption, not in the brief: current LHDN resident rates (YA 2025). "
                   "Replace if Belanjawan 2027 changes the bands.", NOTE)
for col in "ABCD":
    wa[f"{col}{row}"].fill = FLAG_FILL
row += 1
header_row(wa, row, ["Chargeable income above (RM)", "Marginal rate", "Incremental rate", ""])
row += 1
BR_FIRST = row
brackets = [(0, 0), (5000, 0.01), (20000, 0.03), (35000, 0.06), (50000, 0.11), (70000, 0.19),
            (100000, 0.25), (400000, 0.26), (600000, 0.28), (2000000, 0.30)]
for i, (lower, rate) in enumerate(brackets):
    put(wa, f"A{row}", lower, BLUE, RM0)
    put(wa, f"B{row}", rate, BLUE, PCT)
    put(wa, f"C{row}", f"=B{row}" if i == 0 else f"=B{row}-B{row - 1}", BLACK, PCT)
    row += 1
BR_LAST = row - 1
LOWERS = f"Assumptions!$A${BR_FIRST}:$A${BR_LAST}"
INCR = f"Assumptions!$C${BR_FIRST}:$C${BR_LAST}"

row += 1
put(wa, f"A{row}", "Modelling simplifications", BOLD)
for line in [
    "Contributions use the percentage method. Actual EPF, SOCSO and EIS payments follow banded "
    "contribution schedules, so real payslips can differ by a few ringgit.",
    "All statutory bases use gross monthly wage (basic + fixed cash allowances).",
    "PCB is annual tax payable divided by 12 for a single resident with no other income or reliefs.",
    "No bonus, overtime, benefits-in-kind, medical cover or recruitment cost is included.",
]:
    row += 1
    put(wa, f"A{row}", line, NOTE)

wa.column_dimensions["A"].width = 50
wa.column_dimensions["B"].width = 14
wa.column_dimensions["C"].width = 28
wa.column_dimensions["D"].width = 95
wa.freeze_panes = "A5"


# ------------------------------------------------------- formula generators
def er_epf(w):
    return f"={w}*IF({w}<={A['epf_thr']},{A['er_epf_lo']},{A['er_epf_hi']})"


def capped(w, rate, cap):
    return f"=MIN(MIN({w},{A['ceiling']})*{A[rate]},{A[cap]})"


def tax_before_rebate(ci):
    return f"=SUMPRODUCT(({ci}>{LOWERS})*({ci}-{LOWERS})*{INCR})"


def rebate(ci, tax):
    return f"=IF({ci}<={A['rebate_lim']},MIN({tax},{A['rebate']}),0)"


# ---------------------------------------------------- Benchmark Comparison
wc = wb.create_sheet("Benchmark Comparison")
put(wc, "A1", "Benchmark Comparison - Employer Cost and Employee Take-Home", TITLE)
put(wc, "A2", "All amounts in RM per month unless stated. Yellow cell = type your own wage. "
              "Green = linked from Assumptions.", NOTE)
header_row(wc, 3, ["Item (RM)", "National Minimum Wage", "Fresh Graduate Floor",
                   "GLC / GLIC Living Wage", "Custom Scenario"])
COLS = "BCDE"
labels = {
    4: "Basic monthly wage", 5: "Fixed cash allowances", 6: "Gross monthly wage",
    9: "EPF (employer)", 10: "SOCSO (employer)", 11: "EIS (employer)", 12: "HRD Corp levy",
    13: "Total employer on-costs", 14: "Fully loaded monthly cost",
    15: "On-costs as % of gross wage", 16: "Fully loaded annual cost",
    19: "EPF (employee)", 20: "SOCSO (employee)", 21: "EIS (employee)", 22: "PCB / MTD",
    23: "Total employee deductions", 24: "Net take-home pay", 25: "Take-home as % of gross wage",
    28: "Annual gross income", 29: "Less: individual relief", 30: "Less: EPF relief",
    31: "Less: other reliefs", 32: "Chargeable income", 33: "Tax before rebate",
    34: "Less: rebate", 35: "Annual tax payable",
    38: "Fully loaded monthly cost vs floor", 39: "Net take-home pay vs floor",
    40: "Fully loaded annual cost vs floor",
}
for r, label in labels.items():
    put(wc, f"A{r}", label, bold=r in (6, 13, 14, 16, 23, 24, 32, 35))
section(wc, 8, "Employer on-costs", 5)
section(wc, 18, "Employee deductions", 5)
section(wc, 27, "Annual tax computation (basis for PCB)", 5)
section(wc, 37, "Difference vs Fresh Graduate Floor", 5)

for col, src in zip(COLS, ["minwage", "fresh", "glc", None]):
    if src:
        put(wc, f"{col}4", f"={A[src]}", GREEN, RM)
    else:
        put(wc, f"{col}4", 4500, BLUE, RM, fill=YELLOW)
    put(wc, f"{col}5", 0, BLUE, RM)
    w = f"{col}6"
    put(wc, w, f"={col}4+{col}5", fmt=RM, bold=True)
    put(wc, f"{col}9", er_epf(w), fmt=RM)
    put(wc, f"{col}10", capped(w, "er_socso", "er_socso_cap"), fmt=RM)
    put(wc, f"{col}11", capped(w, "er_eis", "er_eis_cap"), fmt=RM)
    put(wc, f"{col}12", f"={w}*{A['hrd']}", fmt=RM)
    put(wc, f"{col}13", f"=SUM({col}9:{col}12)", fmt=RM, bold=True)
    put(wc, f"{col}14", f"={w}+{col}13", fmt=RM, bold=True)
    put(wc, f"{col}15", f"=IF({w}=0,0,{col}13/{w})", fmt=PCT)
    put(wc, f"{col}16", f"={col}14*12", fmt=RM, bold=True)
    put(wc, f"{col}19", f"={w}*{A['ee_epf']}", fmt=RM)
    put(wc, f"{col}20", capped(w, "ee_socso", "ee_socso_cap"), fmt=RM)
    put(wc, f"{col}21", capped(w, "ee_eis", "ee_eis_cap"), fmt=RM)
    put(wc, f"{col}22", f"={col}35/12", fmt=RM)
    put(wc, f"{col}23", f"=SUM({col}19:{col}22)", fmt=RM, bold=True)
    put(wc, f"{col}24", f"={w}-{col}23", fmt=RM, bold=True)
    put(wc, f"{col}25", f"=IF({w}=0,0,{col}24/{w})", fmt=PCT)
    put(wc, f"{col}28", f"={w}*12", fmt=RM)
    put(wc, f"{col}29", f"={A['relief_ind']}", GREEN, RM)
    put(wc, f"{col}30", f"=MIN({col}19*12,{A['relief_epf']})", fmt=RM)
    put(wc, f"{col}31", f"={A['relief_other']}", GREEN, RM)
    put(wc, f"{col}32", f"=MAX(0,{col}28-{col}29-{col}30-{col}31)", fmt=RM, bold=True)
    put(wc, f"{col}33", tax_before_rebate(f"{col}32"), fmt=RM)
    put(wc, f"{col}34", rebate(f"{col}32", f"{col}33"), fmt=RM)
    put(wc, f"{col}35", f"={col}33-{col}34", fmt=RM, bold=True)
    put(wc, f"{col}38", f"={col}14-$C$14", fmt=RM)
    put(wc, f"{col}39", f"={col}24-$C$24", fmt=RM)
    put(wc, f"{col}40", f"={col}16-$C$16", fmt=RM)
    for r in (13, 23, 32, 35):
        wc[f"{col}{r}"].border = TOP

put(wc, "A42", "Allowances default to zero; enter fixed cash allowances in row 5 to include them in "
               "every statutory base.", NOTE)
wc.column_dimensions["A"].width = 38
for col in COLS:
    wc.column_dimensions[col].width = 20
wc.freeze_panes = "B4"

# chart helper block: where each ringgit of loaded cost goes
put(wc, "G3", "Chart data", BOLD)
for i, col in enumerate(COLS):
    r = 4 + i
    put(wc, f"G{r}", f"={col}3")
    put(wc, f"H{r}", f"={col}24", fmt=RM)
    put(wc, f"I{r}", f"={col}23", fmt=RM)
    put(wc, f"J{r}", f"={col}13", fmt=RM)
put(wc, "H3", "Net take-home pay", BOLD)
put(wc, "I3", "Employee deductions", BOLD)
put(wc, "J3", "Employer on-costs", BOLD)
wc.column_dimensions["G"].width = 24
for col in "HIJ":
    wc.column_dimensions[col].width = 20

bar = BarChart()
bar.type = "col"
bar.grouping = "stacked"
bar.overlap = 100
bar.title = "Where the fully loaded monthly cost goes (RM)"
bar.y_axis.title = "RM per month"
bar.y_axis.number_format = "#,##0"
bar.y_axis.delete = False
bar.x_axis.delete = False
bar.add_data(Reference(wc, min_col=8, max_col=10, min_row=3, max_row=7), titles_from_data=True)
bar.set_categories(Reference(wc, min_col=7, min_row=4, max_row=7))
for s, colour in zip(bar.series, ["2E75B6", "A6A6A6", "ED7D31"]):
    s.graphicalProperties.solidFill = colour
    s.graphicalProperties.line.solidFill = colour
bar.legend.position = "b"
bar.height = 9
bar.width = 18
wc.add_chart(bar, "G9")

# ---------------------------------------------------------- Headcount Budget
wh = wb.create_sheet("Headcount Budget")
put(wh, "A1", "Headcount Budget - Annual Cost of a Hiring Plan", TITLE)
put(wh, "A2", "Blue headcounts are example values; overwrite them with your own plan. "
              "Wages and per-head costs link to Benchmark Comparison.", NOTE)
header_row(wh, 4, ["Tier", "Gross monthly wage (RM)", "Headcount",
                   "Loaded cost per head / month (RM)", "Monthly payroll cost (RM)",
                   "Annual payroll cost (RM)", "Annual statutory on-costs (RM)",
                   "Net take-home per head / month (RM)"])
tiers = [("National Minimum Wage", "B", 10), ("Fresh Graduate Floor", "C", 15),
         ("GLC / GLIC Living Wage", "D", 5), ("Custom Scenario", "E", 2)]
for i, (name, col, heads) in enumerate(tiers):
    r = 5 + i
    put(wh, f"A{r}", name)
    put(wh, f"B{r}", f"='Benchmark Comparison'!{col}6", GREEN, RM)
    put(wh, f"C{r}", heads, BLUE, RM0, fill=YELLOW)
    put(wh, f"D{r}", f"='Benchmark Comparison'!{col}14", GREEN, RM)
    put(wh, f"E{r}", f"=C{r}*D{r}", fmt=RM)
    put(wh, f"F{r}", f"=E{r}*12", fmt=RM)
    put(wh, f"G{r}", f"=C{r}*'Benchmark Comparison'!{col}13*12", fmt=RM)
    put(wh, f"H{r}", f"='Benchmark Comparison'!{col}24", GREEN, RM)
put(wh, "A9", "Total", BOLD)
for col in "CEFG":
    put(wh, f"{col}9", f"=SUM({col}5:{col}8)", fmt=RM0 if col == "C" else RM, bold=True)
put(wh, "D9", "=IF(C9=0,0,E9/C9)", fmt=RM, bold=True)
put(wh, "B9", "Average loaded cost per head:", NOTE)
for col in "ABCDEFGH":
    wh[f"{col}9"].border = TOP
put(wh, "A10", "On-costs as % of total annual payroll cost")
put(wh, "G10", "=IF(F9=0,0,G9/F9)", fmt=PCT)

put(wh, "A13", "Wage uplift analysis", BOLD)
put(wh, "A14", "What it costs the employer, and what the employee gains, when staff move up a benchmark.", NOTE)
header_row(wh, 15, ["Uplift", "Wage increase per head / month (RM)", "Staff affected",
                    "Extra loaded cost per head / month (RM)", "Extra monthly cost (RM)",
                    "Extra annual cost (RM)", "Employer cost per RM1 of extra take-home (RM)",
                    "Extra take-home per head / month (RM)"])
uplifts = [("Minimum Wage to Fresh Graduate Floor", "B", "C", 10),
           ("Fresh Graduate Floor to GLC Living Wage", "C", "D", 15),
           ("Minimum Wage to GLC Living Wage", "B", "D", 10)]
for i, (name, lo, hi, heads) in enumerate(uplifts):
    r = 16 + i
    bc = "'Benchmark Comparison'!"
    put(wh, f"A{r}", name)
    put(wh, f"B{r}", f"={bc}{hi}6-{bc}{lo}6", fmt=RM)
    put(wh, f"C{r}", heads, BLUE, RM0, fill=YELLOW)
    put(wh, f"D{r}", f"={bc}{hi}14-{bc}{lo}14", fmt=RM)
    put(wh, f"E{r}", f"=C{r}*D{r}", fmt=RM)
    put(wh, f"F{r}", f"=E{r}*12", fmt=RM)
    put(wh, f"H{r}", f"={bc}{hi}24-{bc}{lo}24", fmt=RM)
    put(wh, f"G{r}", f"=IF(H{r}=0,0,D{r}/H{r})", fmt=RM)
put(wh, "A20", "Each uplift row is independent; the three rows are alternatives, not a sum.", NOTE)
wh.column_dimensions["A"].width = 40
for col in "BCDEFGH":
    wh.column_dimensions[col].width = 21
wh.freeze_panes = "A5"

# --------------------------------------------------------------- Wage Ladder
wl = wb.create_sheet("Wage Ladder")
put(wl, "A1", "Wage Ladder - Sensitivity of Cost and Take-Home to Monthly Wage", TITLE)
put(wl, "A2", "Change the starting wage or step (yellow) to redraw the ladder and chart.", NOTE)
put(wl, "A3", "Starting wage (RM)")
put(wl, "B3", 2000, BLUE, RM, fill=YELLOW)
put(wl, "C3", "Step (RM)")
put(wl, "D3", 250, BLUE, RM, fill=YELLOW)
header_row(wl, 5, ["Gross monthly wage (RM)", "EPF employer", "SOCSO employer", "EIS employer",
                   "HRD levy", "Total on-costs", "Fully loaded cost", "On-cost %",
                   "EPF employee", "SOCSO employee", "EIS employee", "Chargeable income (annual)",
                   "Tax payable (annual)", "PCB / month", "Net take-home pay", "Take-home %"])
L_FIRST, L_LAST = 6, 22
for r in range(L_FIRST, L_LAST + 1):
    w = f"$A{r}"
    put(wl, f"A{r}", "=$B$3" if r == L_FIRST else f"=A{r - 1}+$D$3", fmt=RM, bold=True)
    put(wl, f"B{r}", er_epf(w), fmt=RM)
    put(wl, f"C{r}", capped(w, "er_socso", "er_socso_cap"), fmt=RM)
    put(wl, f"D{r}", capped(w, "er_eis", "er_eis_cap"), fmt=RM)
    put(wl, f"E{r}", f"={w}*{A['hrd']}", fmt=RM)
    put(wl, f"F{r}", f"=SUM(B{r}:E{r})", fmt=RM)
    put(wl, f"G{r}", f"=A{r}+F{r}", fmt=RM, bold=True)
    put(wl, f"H{r}", f"=IF(A{r}=0,0,F{r}/A{r})", fmt=PCT)
    put(wl, f"I{r}", f"={w}*{A['ee_epf']}", fmt=RM)
    put(wl, f"J{r}", capped(w, "ee_socso", "ee_socso_cap"), fmt=RM)
    put(wl, f"K{r}", capped(w, "ee_eis", "ee_eis_cap"), fmt=RM)
    put(wl, f"L{r}", f"=MAX(0,A{r}*12-{A['relief_ind']}-MIN(I{r}*12,{A['relief_epf']})"
                     f"-{A['relief_other']})", fmt=RM)
    tax = tax_before_rebate(f"L{r}")[1:]
    put(wl, f"M{r}", f"={tax}-IF(L{r}<={A['rebate_lim']},MIN({tax},{A['rebate']}),0)", fmt=RM)
    put(wl, f"N{r}", f"=M{r}/12", fmt=RM)
    put(wl, f"O{r}", f"=A{r}-I{r}-J{r}-K{r}-N{r}", fmt=RM, bold=True)
    put(wl, f"P{r}", f"=IF(A{r}=0,0,O{r}/A{r})", fmt=PCT)
wl.column_dimensions["A"].width = 22
for i in range(2, 17):
    wl.column_dimensions[get_column_letter(i)].width = 17
wl.freeze_panes = "B6"

line = LineChart()
line.title = "Fully loaded cost vs net take-home by gross wage (RM / month)"
line.y_axis.title = "RM per month"
line.x_axis.title = "Gross monthly wage (RM)"
line.y_axis.number_format = "#,##0"
line.x_axis.number_format = "#,##0"
line.y_axis.delete = False
line.x_axis.delete = False
for col, colour in ((7, "ED7D31"), (15, "2E75B6")):
    line.add_data(Reference(wl, min_col=col, min_row=5, max_row=L_LAST), titles_from_data=True)
    s = line.series[-1]
    s.graphicalProperties.line.solidFill = colour
    s.graphicalProperties.line.width = 22000
    s.smooth = False
line.set_categories(Reference(wl, min_col=1, min_row=L_FIRST, max_row=L_LAST))
line.legend.position = "b"
line.height = 9
line.width = 22
wl.add_chart(line, "A25")

# ------------------------------------------------------------------ Overview
wo = wb.create_sheet("Overview", 0)
put(wo, "A1", "Belanjawan 2027 Fresh Graduate Wage Cost Model", TITLE)
put(wo, "A2", "Fully loaded employer cost and employee take-home pay at the three Belanjawan 2027 "
              "wage benchmarks.", NOTE)
header_row(wo, 4, ["Headline (RM / month)", "National Minimum Wage", "Fresh Graduate Floor",
                   "GLC / GLIC Living Wage"])
for r, (label, src, fmt) in enumerate([
    ("Gross monthly wage", 6, RM), ("Employer on-costs", 13, RM),
    ("Fully loaded monthly cost", 14, RM), ("On-costs as % of gross wage", 15, PCT),
    ("Net take-home pay", 24, RM), ("PCB / MTD", 22, RM),
    ("Fully loaded annual cost", 16, RM),
], start=5):
    put(wo, f"A{r}", label)
    for col in "BCD":
        put(wo, f"{col}{r}", f"='Benchmark Comparison'!{col}{src}", GREEN, fmt)

put(wo, "A14", "How to use this workbook", BOLD)
guide = [
    ("Assumptions", "Every rate, cap, relief and tax band. Change a blue value and the whole model updates."),
    ("Benchmark Comparison", "Side-by-side payslip and employer cost for each benchmark, plus a custom wage (yellow cell E4)."),
    ("Headcount Budget", "Annual payroll budget for a hiring plan and the cost of moving staff up a benchmark (yellow headcounts)."),
    ("Wage Ladder", "Cost and take-home across a range of wages, with a chart (yellow start and step)."),
]
for i, (sheet, text) in enumerate(guide):
    put(wo, f"A{15 + i}", sheet, BOLD)
    put(wo, f"B{15 + i}", text)

put(wo, "A21", "Colour legend", BOLD)
put(wo, "A22", "Blue text", BLUE)
put(wo, "B22", "Hardcoded input you can edit")
put(wo, "A23", "Black text", BLACK)
put(wo, "B23", "Formula - do not overwrite")
put(wo, "A24", "Green text", GREEN)
put(wo, "B24", "Value linked from another sheet")
put(wo, "A25", "Yellow fill", BLACK, fill=YELLOW)
put(wo, "B25", "Scenario lever: the cells meant to be changed (example values pre-filled)")
put(wo, "A26", "Orange fill", BLACK, fill=FLAG_FILL)
put(wo, "B26", "Assumption to verify against the official schedule before relying on the output")

put(wo, "A28", "Source and limits", BOLD)
put(wo, "A29", "Parameters were supplied by the model author from a Belanjawan 2027 brief and have not "
               "been checked against gazetted schedules. Illustrative only; not payroll or tax advice.", NOTE)
wo.column_dimensions["A"].width = 34
for col in "BCD":
    wo.column_dimensions[col].width = 24

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", OUT)
